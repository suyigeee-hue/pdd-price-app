"""拼多多商品 Excel 批量改价 Streamlit 应用。

启动方式：
    pip install streamlit pandas openpyxl
    streamlit run streamlit_pdd_price_app.py

首次启动会在本脚本同目录自动创建 config.json，用于永久保存价格规则。
"""

from __future__ import annotations

import json
import math
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import load_workbook


st.set_page_config(page_title="拼多多商品批量改价", page_icon="💰", layout="centered")

CONFIG_PATH = Path(__file__).resolve().with_name("config.json")
REQUIRED_COLUMNS = ["商品名称", "规格名称", "单买价", "拼单价"]

# 首次运行时写入本地 config.json 的默认规则，可在页面中自行调整。
DEFAULT_RULES = [
    {"关键词": "10cm", "拼单价": 22.68},
    {"关键词": "20cm", "拼单价": 26.79},
    {"关键词": "30cm", "拼单价": 27.89},
    {"关键词": "40cm", "拼单价": 29.89},
    {"关键词": "鼠标垫", "拼单价": 34.69},
    {"关键词": "老鼠干30cm", "拼单价": 10.75},
    {"关键词": "老鼠干15cm", "拼单价": 9.68},
    {"关键词": "15cm", "拼单价": 9.68},
    {"关键词": "老鼠干", "拼单价": 10.75},
    {"关键词": "7cm猫耳", "拼单价": 24.89},
    {"关键词": "吱吱叫", "拼单价": 23.69},
    {"关键词": "吐泡泡", "拼单价": 24.68},

]


def write_config(rules: list[dict[str, Any]]) -> None:
    """将规则写入本地 JSON；ensure_ascii=False 保证中文可读。"""
    with CONFIG_PATH.open("w", encoding="utf-8") as config_file:
        json.dump({"price_rules": rules}, config_file, ensure_ascii=False, indent=2)


def load_config() -> list[dict[str, Any]]:
    """读取配置；首次启动或配置损坏时使用并保存默认规则。"""
    if not CONFIG_PATH.exists():
        write_config(DEFAULT_RULES)
        return DEFAULT_RULES.copy()

    try:
        with CONFIG_PATH.open("r", encoding="utf-8") as config_file:
            saved_config = json.load(config_file)
        rules = saved_config.get("price_rules", [])
        if not isinstance(rules, list):
            raise ValueError("price_rules 不是列表")
        return rules
    except (json.JSONDecodeError, OSError, ValueError, AttributeError):
        # 配置损坏时不让网页无法启动，恢复为一份可编辑的默认配置。
        write_config(DEFAULT_RULES)
        return DEFAULT_RULES.copy()


def clean_rules(edited_rules: pd.DataFrame) -> list[dict[str, Any]]:
    """校验并清理编辑器中的规则，返回适合保存和匹配的规则列表。"""
    rules: list[dict[str, Any]] = []
    seen_keywords: set[str] = set()

    for row_number, row in edited_rules.iterrows():
        keyword = "" if pd.isna(row["关键词"]) else str(row["关键词"]).strip()

        # data_editor 新增的一整行空白内容不作为规则保存。
        if not keyword and pd.isna(row["拼单价"]):
            continue
        if not keyword:
            raise ValueError(f"第 {row_number + 1} 条规则缺少关键词。")
        if keyword in seen_keywords:
            raise ValueError(f"关键词“{keyword}”重复，请保留一条即可。")

        try:
            group_price = float(row["拼单价"])
        except (TypeError, ValueError):
            raise ValueError(f"关键词“{keyword}”的价格必须是数字。") from None

        if not math.isfinite(group_price):
            raise ValueError(f"关键词“{keyword}”的价格必须是有限数字。")

        seen_keywords.add(keyword)
        rules.append(
            {
                "关键词": keyword,
                "拼单价": group_price,
            }
        )

    if not rules:
        raise ValueError("请至少保留一条有效的价格规则。")
    return rules


def find_matched_rule(product_name: object, spec_name: object, rules: list[dict[str, Any]]) -> dict[str, Any] | None:
    """按关键词长度从长到短匹配，并在命中第一条后立即停止。"""
    product_text = "" if pd.isna(product_name) else str(product_name)
    spec_text = "" if pd.isna(spec_name) else str(spec_name)

    # reverse=True 表示“猫耳10cm”优先于“10cm”。长度相同则保持配置中的原有顺序。
    sorted_rules = sorted(rules, key=lambda rule: len(rule["关键词"]), reverse=True)
    for rule in sorted_rules:
        if rule["关键词"] in product_text or rule["关键词"] in spec_text:
            return rule  # 命中最高优先级规则后立即结束本行匹配。
    return None


def process_excel(source_bytes: bytes, rules: list[dict[str, Any]]) -> tuple[bytes, int]:
    """用 pandas 读取/匹配，只用 openpyxl 回写 E/F，最大程度保留原格式。"""
    # 读取第一个工作表；它应是上传文件中的商品数据表。
    excel_file = pd.ExcelFile(BytesIO(source_bytes))
    sheet_name = excel_file.sheet_names[0]
    dataframe = pd.read_excel(excel_file, sheet_name=sheet_name, dtype=object)

    missing_columns = [column for column in REQUIRED_COLUMNS if column not in dataframe.columns]
    if missing_columns:
        raise ValueError(
            "第一个工作表缺少表头：" + "、".join(missing_columns)
            + "。请确认表头位于第 1 行，且名称与要求一致。"
        )

    # 打开同一份原始工作簿，仅修改第 5、6 列（E/F）第 2 行及之后的单元格。
    # A列商品ID和C列SKUID在此流程中完全不会被写入。
    workbook = load_workbook(BytesIO(source_bytes))
    worksheet = workbook[sheet_name]

    updated_count = 0
    for excel_row, (product_name, spec_name) in enumerate(
        zip(dataframe["商品名称"], dataframe["规格名称"]), start=2
    ):
        matched_rule = find_matched_rule(product_name, spec_name, rules)
        if matched_rule is None:
            continue  # 没有关键词命中时，原 E/F 价格保持不变。

        group_price = matched_rule["拼单价"]
        worksheet.cell(row=excel_row, column=5).value = group_price + 2  # E：单买价 = 拼单价 + 2
        worksheet.cell(row=excel_row, column=6).value = group_price  # F：拼单价 = 配置价格
        updated_count += 1

    result = BytesIO()
    workbook.save(result)
    return result.getvalue(), updated_count


def main() -> None:
    st.title("💰 拼多多商品 Excel 批量改价")
    st.caption("配置价格写入 F 列“拼单价”，E 列“单买价”自动为拼单价 + 2；商品 ID、SKUID 与原有表头样式不会被改动。")

    uploaded_file = st.file_uploader("上传拼多多商品 Excel 文件（.xlsx）", type=["xlsx"])

    st.subheader("价格规则")
    st.caption("可直接增删改。处理时总是优先匹配更长的关键词；相同长度则按当前表格顺序匹配。")
    current_rules = load_config()
    rules_dataframe = pd.DataFrame(current_rules, columns=["关键词", "拼单价"])
    edited_rules = st.data_editor(
        rules_dataframe,
        column_config={
            "关键词": st.column_config.TextColumn("关键词", required=True),
            "拼单价": st.column_config.NumberColumn("拼单价", min_value=0.0, format="%.2f", required=True),
        },
        num_rows="dynamic",
        hide_index=True,
        use_container_width=True,
        key="price_rules_editor",
    )

    if st.button("保存配置", use_container_width=True):
        try:
            valid_rules = clean_rules(edited_rules)
            write_config(valid_rules)
            st.success(f"已保存 {len(valid_rules)} 条规则到 {CONFIG_PATH.name}。")
        except ValueError as error:
            st.error(f"配置未保存：{error}")

    st.divider()
    if st.button("开始处理", type="primary", use_container_width=True):
        if uploaded_file is None:
            st.warning("请先上传一个 .xlsx 文件。")
        else:
            try:
                # 即使尚未点击“保存配置”，处理也会使用编辑器内当前显示的规则。
                valid_rules = clean_rules(edited_rules)
                output_bytes, updated_count = process_excel(uploaded_file.getvalue(), valid_rules)
                st.session_state["output_excel"] = output_bytes
                st.session_state["output_filename"] = f"已改价_{uploaded_file.name}"
                st.success(f"处理完成，共更新 {updated_count} 行。")
            except Exception as error:
                st.error(f"处理失败：{error}")

    if "output_excel" in st.session_state:
        st.download_button(
            "下载处理后的 Excel 文件",
            data=st.session_state["output_excel"],
            file_name=st.session_state["output_filename"],
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )


if __name__ == "__main__":
    main()
